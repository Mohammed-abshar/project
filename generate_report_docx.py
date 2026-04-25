import docx
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, Inches

doc = Document()

# Title
title = doc.add_heading('Final Project Report: Cloud Cost Intelligence Platform for Healthcare', 0)

# Table of Contents
doc.add_heading('Table of Contents', 1)
table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Section'
hdr_cells[1].text = 'Title'
hdr_cells[2].text = 'Page'

toc_data = [
    ('1', 'Executive Summary', '2'),
    ('2', 'Background', '3'),
    ('2.1', 'Aim', '3'),
    ('2.2', 'Technologies', '3'),
    ('2.3', 'Hardware Architecture', '4'),
    ('2.4', 'Software Architecture', '4'),
    ('3', 'System', '5'),
    ('3.1', 'Requirements', '5'),
    ('3.1.1', 'Functional requirements', '5'),
    ('3.1.2', 'User requirements', '5'),
    ('3.1.3', 'Environmental requirements', '6'),
    ('3.2', 'Design and Architecture', '6'),
    ('3.3', 'Implementation', '7'),
    ('3.4', 'Testing', '7'),
    ('3.5', 'Graphical User Interface (GUI) Layout', '8'),
    ('3.6', 'Customer testing', '8'),
    ('3.7', 'Evaluation', '9'),
    ('4', 'Snapshots of the Project', '10'),
    ('5', 'Conclusions', '11'),
    ('6', 'Further development or research', '12'),
    ('7', 'References', '13'),
    ('8', 'Appendix', '14')
]

for section, title, page in toc_data:
    row_cells = table.add_row().cells
    row_cells[0].text = section
    row_cells[1].text = title
    row_cells[2].text = page

doc.add_page_break()

# Executive Summary
doc.add_heading('Executive Summary', 1)
doc.add_paragraph("The Cloud Cost Intelligence Platform for Healthcare is a comprehensive data analytics and visualization system designed to address the growing challenge of managing cloud computing expenses in the healthcare sector. As modern hospitals and medical institutions increasingly adopt multi-cloud infrastructures (AWS, Azure, Google Cloud) for their digital transformation, tracking resource utilization across varied departments—such as Cardiology, Neurology, and Oncology—becomes critically complex and prone to budgetary overruns.")
doc.add_paragraph("This project delivers a robust, interactive Streamlit dashboard that ingests, processes, and visualizes extensive cloud billing datasets. By leveraging Python, Pandas, and Plotly, the platform empowers healthcare IT administrators to monitor monthly cloud costs, analyze department-level spending, and forecast future resource demands. A key feature is its anomaly detection capability, which identifies unexpected cost spikes, ensuring proactive budget management and adherence to financial constraints.")
doc.add_paragraph("Ultimately, this project bridges the gap between raw cloud billing data and actionable intelligence, optimizing operational efficiency without compromising the critical compute and storage resources required for modern healthcare delivery.")

doc.add_page_break()

# Background
doc.add_heading('Background', 1)

doc.add_heading('Aim', 2)
doc.add_paragraph("The primary aim of this project is to develop an intelligent, intuitive, and highly responsive dashboard that visualizes cloud billing data specifically tailored for the healthcare industry. The platform seeks to enable healthcare institutions to effectively monitor their multi-cloud expenditures, track departmental resource utilization, detect billing anomalies, and optimize future resource provisioning based on predictive demand models.")

doc.add_heading('Technologies', 2)
p = doc.add_paragraph()
p.add_run("The project is built upon a modern, Python-centric technology stack:\n").bold = True
doc.add_paragraph("Programming Language: Python 3.10+", style='List Bullet')
doc.add_paragraph("Data Processing: Pandas, NumPy", style='List Bullet')
doc.add_paragraph("Frontend / Dashboarding: Streamlit", style='List Bullet')
doc.add_paragraph("Data Visualization: Plotly, Matplotlib/Seaborn", style='List Bullet')
doc.add_paragraph("Machine Learning / Logic: Scikit-Learn (for anomaly detection and forecasting)", style='List Bullet')
doc.add_paragraph("Data Source: Synthetically generated CSV datasets (healthcare_cloud_billing.csv) emulating real-world multi-cloud billing structures.", style='List Bullet')
doc.add_paragraph("Environment: Jupyter Notebook (prototyping) and standard IDEs (VS Code).", style='List Bullet')

doc.add_heading('Hardware Architecture', 2)
doc.add_paragraph("The hardware architecture is designed to be lightweight, as the current iteration runs primarily in a local or containerized environment before cloud deployment:")
doc.add_paragraph("Client: Any modern web browser (Chrome, Edge, Safari) capable of rendering HTML5/JavaScript.", style='List Bullet')
doc.add_paragraph("Server: A standard multi-core processor (x86_64 architecture) with a minimum of 8GB RAM for processing large Pandas DataFrames efficiently.", style='List Bullet')
doc.add_paragraph("Storage: Minimal local storage required (SSD recommended) for hosting the CSV datasets and application scripts.", style='List Bullet')

doc.add_heading('Software Architecture', 2)
doc.add_paragraph("The software architecture follows a modular, monolithic design centered around data flow:")
doc.add_paragraph("Data Ingestion Layer: Reads raw cloud billing data from CSV formats (generate_dataset.py handles the mock data creation).", style='List Number')
doc.add_paragraph("Data Processing Layer: Utilizes Pandas to clean, aggregate, and compute derived metrics (e.g., usage hours, cost variances, anomaly scores).", style='List Number')
doc.add_paragraph("Application Logic Layer: Contains the core Python functions for filtering data by Hospital, Department, and Date ranges.", style='List Number')
doc.add_paragraph("Presentation Layer: The Streamlit framework renders the processed data into interactive GUI components, including full-width charts, KPI metric cards, and data tables.", style='List Number')

# System
doc.add_heading('System', 1)

doc.add_heading('Requirements', 2)
doc.add_heading('Functional requirements', 3)
doc.add_paragraph("The system must generate and ingest realistic multi-cloud billing data spanning various healthcare departments.", style='List Bullet')
doc.add_paragraph("The system must provide an interactive dashboard with filtering capabilities (by vendor, region, hospital, date).", style='List Bullet')
doc.add_paragraph("The system must calculate and display Key Performance Indicators (KPIs) such as Total Spend, Average Monthly Cost, and Peak Usage.", style='List Bullet')
doc.add_paragraph("The system must generate automated visual reports (Bar charts, Pie charts, Line graphs) for cost breakdown.", style='List Bullet')
doc.add_paragraph("The system must incorporate a predictive mechanism to forecast demand and flag cost anomalies.", style='List Bullet')

doc.add_heading('User requirements', 3)
doc.add_paragraph("The user interface must be intuitive, requiring no technical coding knowledge from the end-user (healthcare administrators).", style='List Bullet')
doc.add_paragraph("The dashboard must utilize a clean, full-width layout for optimal readability of charts and tables.", style='List Bullet')
doc.add_paragraph("Visualizations must be interactive, allowing users to hover over data points for detailed tooltips.", style='List Bullet')
doc.add_paragraph("The system must load quickly and process data filtering requests with minimal latency.", style='List Bullet')

doc.add_heading('Environmental requirements', 3)
doc.add_paragraph("The system must run on standard operating systems (Windows, macOS, Linux).", style='List Bullet')
doc.add_paragraph("The application must be deployable via a standard Python environment (pip install -r requirements.txt).", style='List Bullet')
doc.add_paragraph("No strict dependency on live cloud APIs during the initial deployment phase, ensuring offline analytical capabilities using downloaded datasets.", style='List Bullet')

doc.add_heading('Design and Architecture', 2)
doc.add_paragraph("The platform is designed around a single-page Streamlit application (dashboard.py). The design emphasizes a top-down analytical approach:")
doc.add_paragraph("Header & Global Filters: Located at the top/sidebar to establish the analytical scope.", style='List Number')
doc.add_paragraph("High-Level KPIs: A row of metric cards summarizing the absolute most important figures.", style='List Number')
doc.add_paragraph("Detailed Visualizations: Full-width interactive charts breaking down costs by Service Type (Compute, Storage, AI Services) and Vendor (AWS, Azure, GCP).", style='List Number')
doc.add_paragraph("Tabular Data & Anomalies: Detailed data grids at the bottom for granular inspection and a dedicated section for flagged billing anomalies.", style='List Number')

doc.add_heading('Implementation', 2)
doc.add_paragraph("The implementation relies heavily on Python scripting. generate_dataset.py utilizes numpy and random to populate a DataFrame with synthetic invoices, computing realistic base costs and introducing random noise for anomaly detection. dashboard.py imports this dataset, applies Streamlit's @st.cache_data for performance optimization, and leverages plotly.express to render the UI components sequentially.")

doc.add_heading('Testing', 2)
doc.add_paragraph("Testing was conducted iteratively during development:")
doc.add_paragraph("Unit Testing: Individual data processing functions were tested in Jupyter Notebooks (Healthcare_Cloud_Cost_Intelligence.ipynb) to ensure mathematical accuracy of cost calculations.", style='List Bullet')
doc.add_paragraph("Integration Testing: The Streamlit frontend was tested to ensure it correctly binds to the underlying Pandas data models.", style='List Bullet')
doc.add_paragraph("UI Testing: Visual verification was performed to ensure the layout transitioned successfully from a multi-column format to a clean, full-width presentation.", style='List Bullet')

doc.add_heading('Graphical User Interface (GUI) Layout', 2)
doc.add_paragraph("The GUI is structured into logically separated tabs within the Streamlit app:")
doc.add_paragraph("Overview Tab: Contains KPI metrics and high-level trend lines.", style='List Bullet')
doc.add_paragraph("Departmental Analysis Tab: Focuses on cost distribution among hospitals and specific departments (e.g., Cardiology vs. IT Dept).", style='List Bullet')
doc.add_paragraph("Anomalies & Forecasts Tab: Highlights cost spikes and provides a predictive view of upcoming billing cycles.", style='List Bullet')
doc.add_paragraph("The layout strictly avoids cluttered multi-column designs in favor of expansive, highly readable sequential visualizations.")

doc.add_heading('Customer testing', 2)
doc.add_paragraph("During the final phases, mock customer testing (simulating a healthcare IT Director) was conducted. Feedback indicated that the initial multi-column layout was too dense. Consequently, the UI was refactored to a full-width sequential layout, significantly improving the user's ability to digest complex financial graphs.")

doc.add_heading('Evaluation', 2)
doc.add_paragraph("The project successfully meets its core objectives. It provides a highly responsive, accurate, and visually compelling analysis of cloud costs. The synthetic data generation is robust enough to mimic real-world unpredictability, while the Streamlit dashboard effectively translates this complexity into actionable insights.")

doc.add_heading('Snapshots of the Project', 1)
doc.add_paragraph("(Please insert your dashboard screenshots here. Recommend showing the main KPIs, the full-width bar charts, and the Anomaly detection tab.)")

doc.add_heading('Conclusions', 1)
doc.add_paragraph("The 'Cloud Cost Intelligence Platform for Healthcare' demonstrates the immense value of applying data analytics to cloud infrastructure management. As healthcare providers continue to migrate sensitive workloads (like AI diagnostic services and patient databases) to the cloud, tools that provide immediate financial visibility become indispensable. By successfully building a complete pipeline—from realistic data generation to interactive visualization—this project proves that cost transparency can be achieved through accessible Python frameworks. The transition to a clean, full-width UI further underscores the importance of user-centric design in data-heavy enterprise applications.")

doc.add_heading('Further development or research', 1)
doc.add_paragraph("Future enhancements for this platform could include:")
doc.add_paragraph("Live API Integration: Replacing the synthetic CSV data with live API connections to AWS Cost Explorer, Azure Billing API, and GCP Billing Export.", style='List Number')
doc.add_paragraph("Advanced Machine Learning: Implementing sophisticated deep learning models (e.g., LSTMs) for more accurate long-term cost forecasting based on seasonal healthcare trends.", style='List Number')
doc.add_paragraph("HIPAA Compliance Tracking: Integrating a specific compliance scoring metric to ensure that cloud resource utilization aligns with stringent healthcare data privacy regulations.", style='List Number')
doc.add_paragraph("Automated Alerting: Developing a module to send automated email or Slack notifications when daily spending exceeds a predefined threshold.", style='List Number')

doc.add_heading('References', 1)
doc.add_paragraph("1. Python Software Foundation. (2023). Python Language Reference. Available at http://www.python.org")
doc.add_paragraph("2. McKinney, W. (2010). Data Structures for Statistical Computing in Python. Proceedings of the 9th Python in Science Conference, 51-56.")
doc.add_paragraph("3. Streamlit Inc. (2023). Streamlit Documentation. Available at https://docs.streamlit.io/")
doc.add_paragraph("4. Plotly Technologies Inc. (2015). Collaborative data science. Plotly. Available at https://plot.ly")
doc.add_paragraph("5. Cloud Computing in Healthcare: Trends and Security. (2022). Healthcare IT News.")

doc.add_heading('Appendix', 1)
doc.add_heading('Appendix A: Data Generation Logic snippet', 2)
doc.add_paragraph("base_cost = usage_hours * 0.15 + storage_used * 0.02\nif service == 'AI Services':\n    base_cost *= 2.5\nelif service == 'Database':\n    base_cost *= 1.5")

doc.add_heading('Appendix B: Installation Instructions', 2)
doc.add_paragraph("pip install -r requirements.txt\nstreamlit run dashboard.py")

doc.save('Final_Project_Report.docx')
print('Final_Project_Report.docx created successfully.')
